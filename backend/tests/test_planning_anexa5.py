# Planning calendar (appointments CRUD/conflicts/status->WO/capacity/reminder cron) + QC Anexa 5 rule + Anexa 5 register/exports.
# Run serially: pytest /app/backend/tests/test_planning_anexa5.py -n 0 -v   (module fixture creates TEST data and cleans everything)
import csv
import io
import os
import secrets
import sys
import time
from datetime import datetime, timedelta, timezone

import fitz
import openpyxl
import pytest
import requests
from bson import ObjectId

sys.path.insert(0, "/app/backend/tests")
from storno_fixtures import BASE, creds, mdb  # noqa: E402

SIG = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
QC_KEYS = ["noxe", "frane", "geometrie", "faruri_geamuri_parbrize", "parametri_reglaje", "etanșeitate", "proba_drum"]
TODAY = datetime.now(timezone.utc).date().isoformat()


def login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE}/auth/login", json={"email": email, "password": password})
    r.raise_for_status()
    s.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return s


def demo(email):
    return login(email, "Demo2026!")


def z(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%S.000Z")


@pytest.fixture(scope="module")
def ctx():
    db = mdb()
    tag = secrets.token_hex(3).upper()
    st = {"tag": tag, "counters": {c: (db.counters.find_one({"_id": c}) or {}).get("seq") for c in ("work_order", "invoice")},
          "notif_before": {str(n["_id"]) for n in db.notifications.find({}, {"_id": 1})},
          "cron_ids": [], "appts": [], "wos": [], "invoices": []}
    cfg = db.efactura_config.find_one({"_id": "main"}) or {}
    st["auto_send"] = cfg.get("auto_send", False)
    db.efactura_config.update_one({"_id": "main"}, {"$set": {"auto_send": False}})
    now = datetime.now(timezone.utc).isoformat()
    cid = db.clients.insert_one({"name": f"TEST_Client {tag}", "type": "pf", "phone": "0700000000", "email": f"test_{tag.lower()}@example.com", "created_at": now}).inserted_id
    cid2 = db.clients.insert_one({"name": f"TEST_Client2 {tag}", "type": "pf", "phone": "0700000001", "created_at": now}).inserted_id
    vid = db.vehicles.insert_one({"client_id": str(cid), "vin": f"TESTVIN{tag}".ljust(17, "0"), "plate": f"TST{tag}", "make": "Test", "model": "Car",
                                  "mileage": 1000, "created_at": now}).inserted_id
    vid2 = db.vehicles.insert_one({"client_id": str(cid2), "vin": f"TESTVJN{tag}".ljust(17, "0"), "plate": f"TSU{tag}", "make": "Test", "model": "Two",
                                   "mileage": 2000, "created_at": now}).inserted_id
    st.update(client=str(cid), client2=str(cid2), vehicle=str(vid), vehicle2=str(vid2))
    st["admin"] = login(*creds())
    st["sef"] = demo("sef.service@serviceauto.ro")
    st["elevators"] = (db.settings.find_one({"_id": "company"}) or {}).get("elevators", [])
    yield st
    # ---- teardown
    db.appointments.delete_many({"_id": {"$in": [ObjectId(a) for a in st["appts"]]}})
    db.appointments.delete_many({"client_id": {"$in": [st["client"], st["client2"]]}})
    wo_ids = [str(w["_id"]) for w in db.work_orders.find({"client_id": {"$in": [st["client"], st["client2"]]}}, {"_id": 1})]
    inv_ids = [str(i["_id"]) for i in db.invoices.find({"work_order_id": {"$in": wo_ids}}, {"_id": 1})]
    db.invoices.delete_many({"_id": {"$in": [ObjectId(i) for i in inv_ids]}})
    db.payments.delete_many({"invoice_id": {"$in": inv_ids}})
    db.efactura_logs.delete_many({"invoice_id": {"$in": inv_ids}})
    db.rar_autopass_outbox.delete_many({"work_order_id": {"$in": wo_ids}})
    db.time_entries.delete_many({"work_order_id": {"$in": wo_ids}})
    db.work_orders.delete_many({"_id": {"$in": [ObjectId(w) for w in wo_ids]}})
    db.notifications.delete_many({"_id": {"$nin": [ObjectId(n) for n in st["notif_before"]]}, "kind": {"$in": ["appointment_confirmation", "appointment_reminder"]}})
    db.cron_runs.delete_many({"_id": {"$in": st["cron_ids"]}})
    db.vehicles.delete_many({"_id": {"$in": [vid, vid2]}})
    db.clients.delete_many({"_id": {"$in": [cid, cid2]}})
    db.audit_logs.delete_many({"$or": [{"entity_id": {"$in": wo_ids + inv_ids}}, {"details": {"$regex": tag}}]})
    for c, v in st["counters"].items():
        if v is not None:
            db.counters.update_one({"_id": c}, {"$set": {"seq": v}})
    db.efactura_config.update_one({"_id": "main"}, {"$set": {"auto_send": st["auto_send"]}})
    print("RESIDUE", db.clients.count_documents({"name": {"$regex": tag}}), db.appointments.count_documents({"client_id": st["client"]}))


def appt(ctx, start, hours=2, **kw):
    body = {"client_id": ctx["client"], "vehicle_id": ctx["vehicle"], "starts_at": z(start), "ends_at": z(start + timedelta(hours=hours)),
            "elevator": "", "mechanic_ids": [], "advisor": "TEST", "title": f"TEST_Revizie {ctx['tag']}", "notes": ""} | kw
    return body


BASE_DT = datetime(2031, 3, 10, 8, 0, tzinfo=timezone.utc)


# --- RBAC / interval validation
class TestA_Access:
    def test_rbac(self, ctx):
        q = {"start": z(BASE_DT), "end": z(BASE_DT + timedelta(days=1))}
        assert requests.get(f"{BASE}/appointments", params=q).status_code == 401
        assert demo("mecanic@serviceauto.ro").get(f"{BASE}/appointments", params=q).status_code == 403
        assert demo("director@serviceauto.ro").get(f"{BASE}/appointments", params=q).status_code == 200
        assert demo("director@serviceauto.ro").get(f"{BASE}/reports/final-inspections", params={"start": TODAY, "end": TODAY}).status_code == 200
        assert demo("mecanic@serviceauto.ro").get(f"{BASE}/reports/final-inspections", params={"start": TODAY, "end": TODAY}).status_code == 403

    def test_inverse_interval_rejected(self, ctx):
        s = ctx["admin"]
        q = {"start": z(BASE_DT + timedelta(days=1)), "end": z(BASE_DT)}
        assert s.get(f"{BASE}/appointments", params=q).status_code == 400
        assert s.get(f"{BASE}/appointments/capacity", params=q).status_code == 400
        assert s.post(f"{BASE}/appointments", json=appt(ctx, BASE_DT, hours=-1)).status_code == 400

    def test_garbage_datetime_rejected(self, ctx):
        r = ctx["admin"].post(f"{BASE}/appointments", json=appt(ctx, BASE_DT) | {"starts_at": "abc", "ends_at": "abd"})
        if r.status_code == 200:
            ctx["appts"].append(r.json()["id"])
        assert r.status_code in (400, 422), f"non-ISO datetimes accepted: {r.status_code}"

    def test_vehicle_must_belong_to_client(self, ctx):
        r = ctx["admin"].post(f"{BASE}/appointments", json=appt(ctx, BASE_DT) | {"vehicle_id": ctx["vehicle2"]})
        assert r.status_code == 400


# --- CRUD + conflicts
class TestB_Crud:
    def test_create_and_list(self, ctx):
        s = ctx["admin"]
        el = ctx["elevators"][0]
        r = s.post(f"{BASE}/appointments", json=appt(ctx, BASE_DT, elevator=el, mechanic_ids=["TEST_M1"]))
        assert r.status_code == 200, r.text
        a = r.json()
        ctx["appts"].append(a["id"]); ctx["a1"] = a["id"]
        assert a["status"] == "programat" and a["client_name"].startswith("TEST_Client") and a["vehicle_label"].startswith("TST")
        assert "_id" not in a
        lst = s.get(f"{BASE}/appointments", params={"start": z(BASE_DT - timedelta(hours=1)), "end": z(BASE_DT + timedelta(days=1))}).json()
        assert any(x["id"] == a["id"] for x in lst)
        n = mdb().notifications.find_one({"kind": "appointment_confirmation", "appointment_id": a["id"]})
        assert n and n["status"] == "simulated"

    def test_elevator_conflict(self, ctx):
        r = ctx["admin"].post(f"{BASE}/appointments", json=appt(ctx, BASE_DT + timedelta(hours=1), elevator=ctx["elevators"][0]))
        if r.status_code == 200:
            ctx["appts"].append(r.json()["id"])
        assert r.status_code == 409

    def test_mechanic_conflict(self, ctx):
        r = ctx["admin"].post(f"{BASE}/appointments", json=appt(ctx, BASE_DT + timedelta(hours=1), elevator=ctx["elevators"][1], mechanic_ids=["TEST_M1"]))
        if r.status_code == 200:
            ctx["appts"].append(r.json()["id"])
        assert r.status_code == 409

    def test_adjacent_ok_and_edit_move(self, ctx):
        s = ctx["admin"]
        r = s.post(f"{BASE}/appointments", json=appt(ctx, BASE_DT + timedelta(hours=2), elevator=ctx["elevators"][0]))
        assert r.status_code == 200, r.text
        ctx["appts"].append(r.json()["id"]); ctx["a2"] = r.json()["id"]
        # move a2 onto a1 -> conflict
        assert s.put(f"{BASE}/appointments/{ctx['a2']}", json=appt(ctx, BASE_DT + timedelta(hours=1), elevator=ctx["elevators"][0])).status_code == 409
        # move a2 to next day + edit title
        body = appt(ctx, BASE_DT + timedelta(days=1), elevator=ctx["elevators"][0]) | {"title": f"TEST_Edit {ctx['tag']}"}
        r = s.put(f"{BASE}/appointments/{ctx['a2']}", json=body)
        assert r.status_code == 200 and r.json()["title"].startswith("TEST_Edit") and r.json()["starts_at"] == body["starts_at"]

    def test_update_vehicle_ownership_checked(self, ctx):
        r = ctx["admin"].put(f"{BASE}/appointments/{ctx['a2']}", json=appt(ctx, BASE_DT + timedelta(days=1)) | {"vehicle_id": ctx["vehicle2"]})
        assert r.status_code == 400, f"PUT accepts vehicle of other client: {r.status_code}"

    def test_capacity(self, ctx):
        s = ctx["admin"]
        r = s.get(f"{BASE}/appointments/capacity", params={"start": z(BASE_DT.replace(hour=0)), "end": z(BASE_DT.replace(hour=0) + timedelta(days=1))})
        assert r.status_code == 200
        d = r.json()
        assert d["elevators"] == ctx["elevators"] and d["count"] == 1 and isinstance(d["occupancy"], int)
        # a1 is 2h on 1 elevator of N over 8h day => expected 2/(N*8)
        exp = round(120 / (len(ctx["elevators"]) * 480) * 100)
        ctx["cap_actual"] = d["occupancy"]
        assert d["occupancy"] == exp, f"occupancy ignores duration: got {d['occupancy']} expected {exp}"


# --- status -> work order once
class TestC_Status:
    def test_in_lucru_converts_once(self, ctx):
        s = ctx["admin"]
        assert s.post(f"{BASE}/appointments/{ctx['a1']}/status", json={"status": "bogus"}).status_code == 400
        r1 = s.post(f"{BASE}/appointments/{ctx['a1']}/status", json={"status": "in_lucru"})
        assert r1.status_code == 200
        wo = r1.json()["work_order_id"]
        assert wo
        ctx["wos"].append(wo)
        s.post(f"{BASE}/appointments/{ctx['a1']}/status", json={"status": "sosit"})
        r2 = s.post(f"{BASE}/appointments/{ctx['a1']}/status", json={"status": "in_lucru"})
        assert r2.json()["work_order_id"] == wo
        assert mdb().work_orders.count_documents({"client_id": ctx["client"]}) == 1
        w = s.get(f"{BASE}/work-orders/{wo}").json()
        assert w["status"] == "in_lucru" and w["vehicle_id"] == ctx["vehicle"] and w["elevator"] == ctx["elevators"][0]

    def test_parallel_in_lucru_no_duplicate(self, ctx):
        import concurrent.futures as cf
        s = ctx["admin"]
        a = s.post(f"{BASE}/appointments", json=appt(ctx, BASE_DT + timedelta(days=3))).json()
        ctx["appts"].append(a["id"])
        with cf.ThreadPoolExecutor(4) as ex:
            list(ex.map(lambda _: s.post(f"{BASE}/appointments/{a['id']}/status", json={"status": "in_lucru"}), range(4)))
        n = mdb().work_orders.count_documents({"client_id": ctx["client"]})
        assert n == 2, f"race creates duplicate WOs: {n - 1} WOs for one appointment"


# --- reminder cron
class TestD_Cron:
    def test_auth(self, ctx):
        assert requests.post(f"{BASE}/cron/appointments/reminders").status_code == 401
        assert requests.post(f"{BASE}/cron/appointments/reminders", headers={"Authorization": "Bearer wrong"}).status_code == 401

    def test_reminder_idempotent(self, ctx):
        from dotenv import dotenv_values
        secret = dotenv_values("/app/backend/.env")["WEBHOOK_CRON_SECRET"]
        start = datetime.now(timezone.utc) + timedelta(hours=24)
        a = ctx["admin"].post(f"{BASE}/appointments", json=appt(ctx, start, hours=1)).json()
        ctx["appts"].append(a["id"])
        rid = f"TEST_appt_rem_{ctx['tag']}"
        ctx["cron_ids"] += [rid, rid + "_2"]
        h = {"Authorization": f"Bearer {secret}", "X-Webhook-Id": rid}
        r = requests.post(f"{BASE}/cron/appointments/reminders", headers=h)
        assert r.status_code == 200 and r.json() == {"ok": True}
        r = requests.post(f"{BASE}/cron/appointments/reminders", headers=h)
        assert r.status_code == 200 and r.json().get("duplicate") is True
        time.sleep(2)
        requests.post(f"{BASE}/cron/appointments/reminders", headers=h | {"X-Webhook-Id": rid + "_2"})
        time.sleep(2)
        rows = list(mdb().notifications.find({"kind": "appointment_reminder", "appointment_id": a["id"]}))
        assert len(rows) == 1 and rows[0]["status"] == "simulated"


def make_wo(ctx, status="in_lucru"):
    db = mdb()
    lines = [{"id": "l_lab", "type": "labor", "description": "TEST Manopera", "qty": 1, "unit_price": 100, "discount": 0, "vat_rate": 21,
              "product_id": "", "code": "", "unit": "ore"}]
    wid = db.work_orders.insert_one({"number": f"TEST-WO-{ctx['tag']}-{secrets.token_hex(2)}", "client_id": ctx["client"], "vehicle_id": ctx["vehicle"],
                                     "status": status, "mileage_in": 1000, "lines": lines, "quotes": [], "qc": {}, "status_history": [],
                                     "created_at": datetime.now(timezone.utc).isoformat(), "reception": {}, "mechanic_ids": []}).inserted_id
    return str(wid)


def qc_body(code="C", **over):
    items = {k: code for k in QC_KEYS} | over
    return {"items": items, "notes": "TEST", "mechanic_signature": SIG, "chief_signature": SIG}


# --- QC rule
class TestE_QC:
    def test_mechanic_cannot_chief_sign(self, ctx):
        ctx["wo_qc"] = make_wo(ctx)
        r = demo("mecanic@serviceauto.ro").put(f"{BASE}/work-orders/{ctx['wo_qc']}/qc", json=qc_body())
        assert r.status_code == 403

    def test_n_fails_and_blocks_finalize(self, ctx):
        s = ctx["sef"]
        r = s.put(f"{BASE}/work-orders/{ctx['wo_qc']}/qc", json=qc_body(frane="N"))
        assert r.status_code == 200 and r.json()["qc"]["passed"] is False
        assert s.post(f"{BASE}/work-orders/{ctx['wo_qc']}/status", json={"status": "finalizata"}).status_code == 400
        assert ctx["admin"].post(f"{BASE}/invoices/from-work-order/{ctx['wo_qc']}", json={"type": "fiscal"}).status_code == 400

    def test_missing_signature_or_item_not_passed(self, ctx):
        s = ctx["sef"]
        b = qc_body(); b["mechanic_signature"] = ""
        assert s.put(f"{BASE}/work-orders/{ctx['wo_qc']}/qc", json=b).json()["qc"]["passed"] is False
        b = qc_body(); del b["items"]["proba_drum"]
        assert s.put(f"{BASE}/work-orders/{ctx['wo_qc']}/qc", json=b).json()["qc"]["passed"] is False

    def test_conform_passes_and_finalize(self, ctx):
        s = ctx["sef"]
        r = s.put(f"{BASE}/work-orders/{ctx['wo_qc']}/qc", json=qc_body(noxe="NA"))
        q = r.json()["qc"]
        assert q["passed"] is True and q["inspector_name"] == "Șef Service Demo"
        assert s.post(f"{BASE}/work-orders/{ctx['wo_qc']}/status", json={"status": "finalizata"}).status_code == 200

    def test_qc_downgrade_after_finalize_blocks_invoice(self, ctx):
        wid = make_wo(ctx)
        s = ctx["sef"]
        s.put(f"{BASE}/work-orders/{wid}/qc", json=qc_body())
        assert s.post(f"{BASE}/work-orders/{wid}/status", json={"status": "finalizata"}).status_code == 200
        r = s.put(f"{BASE}/work-orders/{wid}/qc", json=qc_body(frane="N"))
        assert r.json()["qc"]["passed"] is False
        inv = ctx["admin"].post(f"{BASE}/invoices/from-work-order/{wid}", json={"type": "fiscal"})
        ctx["downgrade_invoice_status"] = inv.status_code
        assert inv.status_code == 400, "fiscal invoice issued although current QC is N (status stays finalizata after QC downgrade)"


# --- Anexa 5 register
class TestF_Anexa5:
    def test_validation(self, ctx):
        s = ctx["admin"]
        assert s.get(f"{BASE}/reports/final-inspections").status_code == 422
        assert s.get(f"{BASE}/reports/final-inspections", params={"start": "2026-02-01", "end": "2026-01-01"}).status_code == 400
        assert s.get(f"{BASE}/reports/final-inspections/pdf", params={"start": "2026-02-01", "end": "2026-01-01"}).status_code == 400
        assert s.get(f"{BASE}/reports/final-inspections/doc", params={"start": TODAY, "end": TODAY}).status_code == 404

    def test_rows_13_columns_and_inspector_filter(self, ctx):
        s = ctx["admin"]
        d = s.get(f"{BASE}/reports/final-inspections", params={"start": TODAY, "end": TODAY}).json()
        mine = [r for r in d["rows"] if r["vehicle"] == f"TST{ctx['tag']}"]
        assert mine, "TEST QC rows missing"
        row = mine[0]
        assert len(row["columns"]) == 13
        assert row["noxe"] in ("C", "N", "NA") and "Șef Service Demo" in row["inspector"]
        assert all(row[k] in ("C", "N", "NA") for k in ["noxe", "frane", "geometrie", "faruri", "parametri", "etanseitate", "proba"])
        assert "Șef Service Demo" in s.get(f"{BASE}/reports/final-inspections/inspectors").json()
        f = s.get(f"{BASE}/reports/final-inspections", params={"start": TODAY, "end": TODAY, "inspector": "Șef Service Demo"}).json()["rows"]
        assert f and all("Șef Service Demo" in r["inspector"] for r in f)
        none = s.get(f"{BASE}/reports/final-inspections", params={"start": TODAY, "end": TODAY, "inspector": "TEST_Nobody"}).json()["rows"]
        assert none == []
        assert s.get(f"{BASE}/reports/final-inspections", params={"start": "2001-01-01", "end": "2001-01-02"}).json()["rows"] == []

    def test_pdf_a4_landscape_diacritics(self, ctx):
        r = ctx["admin"].get(f"{BASE}/reports/final-inspections/pdf", params={"start": TODAY, "end": TODAY})
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/pdf") and r.content[:4] == b"%PDF"
        open("/tmp/anexa5_test.pdf", "wb").write(r.content)
        doc = fitz.open(stream=r.content, filetype="pdf")
        p = doc[0]
        assert round(p.rect.width) == 842 and round(p.rect.height) == 595
        txt = p.get_text()
        assert "EVIDENȚĂ" in txt and "Frâne" in txt and "Etanșeitate" in txt and f"TST{ctx['tag']}" in txt

    def test_csv_xlsx(self, ctx):
        s = ctx["admin"]
        c = s.get(f"{BASE}/reports/final-inspections/csv", params={"start": TODAY, "end": TODAY})
        assert c.status_code == 200 and c.content[:3] == b"\xef\xbb\xbf"
        rows = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
        assert len(rows[0]) == 13 and rows[0][5] == "Frâne"
        x = s.get(f"{BASE}/reports/final-inspections/xlsx", params={"start": TODAY, "end": TODAY})
        assert x.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        ws = wb.active
        hdr = [c.value for c in ws[1]]
        assert len(hdr) == 13 and hdr[9] == "Etanșeitate"
        cell = ws["A1"]
        ctx["xlsx_header"] = (cell.font.color.rgb if cell.font.color else None, cell.fill.fgColor.rgb if cell.fill else None, cell.fill.fill_type)
        assert not (cell.font.color and cell.font.color.rgb.endswith("FFFFFF") and cell.fill.fill_type is None), \
            f"XLSX header white text on no fill (invisible): {ctx['xlsx_header']}"

    def test_export_content_disposition(self, ctx):
        r = ctx["admin"].get(f"{BASE}/reports/final-inspections/csv", params={"start": TODAY, "end": TODAY})
        assert "attachment" in r.headers.get("content-disposition", ""), "no Content-Disposition filename on exports"
